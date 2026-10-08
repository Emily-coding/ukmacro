"""ONS Business Insights and Conditions Survey (BICS): harvesting the wave files.

BICS is published as one spreadsheet per fortnightly wave. Each one contains
"TS (WTD)" sheets: weighted time series for every question asked in that wave,
carrying that question's history back to when it was first asked. So the latest
copy of each question's sheet is the whole series, and harvesting is mostly
"keep the newest version of every sheet you have seen".

Questions are identified by (normalised question text, population, weighting),
not by sheet name, because ONS renames sheets ("TS1", "TS2", "(2)") but the question
text is stable once months and years are stripped. Weighting is "count" (each
business counts once) or "employment" (businesses weighted by their headcount).

Only the files that are needed get downloaded: the newest wave (for its question log,
which records every question ever asked and in which waves), plus, for each question
the curated series use, the newest wave that asked it.

Stored under data/bics/:
    harvest/<id>.csv   wave, breakdown, then one column per answer (%), one file per question
    catalogue.csv      one row per harvested question: text, population, weighting, answers, waves
    question_log.csv   every question ever asked (from the newest file's logs) and its waves
    waves.csv          wave -> reference period and survey live period
    state.json         which wave files have been harvested
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl

from .http import get, get_json

SITE = "https://www.ons.gov.uk"
DATASET = "/economy/economicoutputandproductivity/output/datasets/businessinsightsandimpactontheukeconomy"

MONTH = r"(January|February|March|April|May|June|July|August|September|October|November|December)"
SIZE_BANDS = re.compile(r"^(\d|all size bands)", re.I)


# ---------------------------------------------------------------- download

def list_waves() -> dict[int, str]:
    """{wave number: dataset page uri} for every published wave."""
    idx = get_json(f"{SITE}{DATASET}/data")
    return {int(d["uri"].rsplit("wave", 1)[1]): d["uri"] for d in idx["datasets"]}


def download_wave(uri: str, dest: Path) -> Path:
    """Download one wave's spreadsheet. File names vary by era, so read them from the page."""
    fname = get_json(f"{SITE}{uri}/data")["downloads"][0]["file"]
    dest.write_bytes(get(f"{SITE}/file?uri={uri}/{fname}", timeout=180))
    return dest


# ---------------------------------------------------------------- normalising

def norm(text: str) -> str:
    """Lower-case, straighten quotes, collapse whitespace, drop trailing '?'.
    Used for matching only; stored text keeps its original form."""
    t = str(text).replace("’", "'").replace("‘", "'").replace(" ", " ")
    t = re.sub(r"\s+", " ", t).strip().rstrip("?").strip().lower()
    return t


def generic_question(text: str) -> str:
    """Strip the 'Question:' prefix and replace concrete months/years with
    placeholders, so 'in July 2026' and 'in [month] [year]' are the same question."""
    t = re.sub(r"^\s*question:\s*", "", str(text), flags=re.I)
    t = re.sub(rf"\b{MONTH}\s+\d{{4}}\b", "[month] [year]", t)
    t = re.sub(rf"\b{MONTH}\b", "[month]", t)
    t = re.sub(r"\b(19|20)\d{2}\b", "[year]", t)
    return re.sub(r"\s+", " ", t).strip()


def parse_base(text: str) -> tuple[str, str]:
    """'As a percentage of businesses ..., broken down by ..., weighted by count, UK, dates'
    -> (population, weighting)."""
    t = str(text or "")
    m = re.search(r"weighted by (\w+)", t, re.I)
    weighting = m.group(1).lower() if m else "count"
    population = re.split(r",\s*(broken down|weighted|UK\b)", t, maxsplit=1)[0].strip()
    return population, weighting


def question_id(question: str, population: str, weighting: str) -> str:
    """Short stable id for the (question, population, weighting) triple; used as file name."""
    words = re.findall(r"[a-z]+", norm(question))
    stop = {"the", "of", "a", "your", "business", "business's", "s", "in", "to", "if", "any",
            "which", "following", "what", "how", "does", "did", "is", "has", "have", "for", "or"}
    slug = "-".join([w for w in words if w not in stop][:6])
    h = hashlib.sha1(f"{norm(question)}|{norm(population)}|{weighting}".encode()).hexdigest()[:6]
    return f"{slug}-{weighting[:3]}-{h}"


def breakdown_group(b: str) -> str:
    """'total', 'size' (e.g. '10 - 49', 'All size bands excluding 0 - 9') or 'industry'."""
    if b == "All businesses":
        return "total"
    return "size" if SIZE_BANDS.match(b) else "industry"


def _num(v) -> float | None:
    """Cells are proportions (0-1). Suppressed / not-applicable markers ('[c]', '[x]') -> None."""
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def _wave_no(v) -> int | None:
    """'Wave 163' -> 163."""
    m = re.search(r"(\d+)", str(v or ""))
    return int(m.group(1)) if m else None


def _parse_period(text: str) -> tuple[str, str] | tuple[None, None]:
    """'1 July 2026 to 31 July 2026' -> ('2026-07-01', '2026-07-31')."""
    parts = re.split(r"\s+to\s+", str(text or "").strip())
    try:
        d = [datetime.strptime(p.strip(), "%d %B %Y").date().isoformat() for p in parts]
        return d[0], d[-1]
    except ValueError:
        return None, None


# ---------------------------------------------------------------- reading a wave file

def _is_header(r) -> bool:
    """The table header row: 'Dates' (or 'Date') | 'Wave' | breakdown | answers..."""
    return bool(r) and len(r) > 2 and str(r[0] or "").strip() in ("Dates", "Date") \
        and str(r[1] or "").strip() == "Wave"


def read_wave(path: Path) -> tuple[list[dict], dict[int, dict]]:
    """Return (tables, periods) from one wave spreadsheet.

    tables:  [{question, population, weighting, sheet, answers, rows: [(wave, breakdown, answer, pct)]}]
    periods: {wave: {ref_start, ref_end, live_start, live_end}} from the 'Dates' sheet
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tables, periods = [], {}

    if "Dates" in wb.sheetnames:
        for r in wb["Dates"].iter_rows(values_only=True):
            w = _wave_no(r[0]) if r and str(r[0] or "").startswith("Wave") else None
            if w is not None and len(r) >= 3:
                rs, re_ = _parse_period(r[1])
                ls, le = _parse_period(r[2])
                periods[w] = {"ref_start": rs, "ref_end": re_, "live_start": ls, "live_end": le}

    for name in wb.sheetnames:
        if "TS" not in name:
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        qrow = next((r[0] for r in rows[:6] if r and str(r[0] or "").lower().startswith("question")), None)
        hi = next((i for i, r in enumerate(rows) if _is_header(r)), None)
        if qrow is None or hi is None:
            continue  # not a standard weighted time-series table
        population, weighting = parse_base(rows[1][0] if len(rows) > 1 else "")
        header, seen = [], {}
        for h in rows[hi]:
            h = re.sub(r"\s+", " ", str(h)).strip() if h is not None else None
            if h:  # ONS occasionally repeats an answer label in one table; keep both
                seen[h] = seen.get(h, 0) + 1
                h = h if seen[h] == 1 else f"{h} ({seen[h]})"
            header.append(h)
        answers = [h for h in header[3:] if h]
        out = []
        for r in rows[hi + 1:]:
            if not r or r[1] is None:
                continue
            w = _wave_no(r[1])
            b = re.sub(r"\s+", " ", str(r[2] or "")).strip()
            b = b[:1].upper() + b[1:].lower()  # older files Title-Case some names ("All Businesses")
            for j, ans in enumerate(header[3:], start=3):
                v = _num(r[j]) if ans and j < len(r) else None
                if v is not None:
                    out.append((w, b, ans, round(v * 100, 2)))
        if out:
            tables.append({"question": generic_question(qrow), "population": population,
                           "weighting": weighting, "sheet": name, "answers": answers, "rows": out})
    return tables, periods


# ---------------------------------------------------------------- the question log

def read_question_log(path: Path) -> list[dict]:
    """Every question ever asked, from the 'Log' sheets of a wave file (the newest file
    has all of them): [{section, subsection, question, rotation, waves: [int]}]."""
    wb = openpyxl.load_workbook(path, read_only=True)
    out = []
    for name in wb.sheetnames:
        if "Log" not in name or name == "Question Logs":
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rows)
                   if r and "Question" in [str(c).strip() for c in r if c]), None)
        if hi is None:
            continue
        hdr = [str(c or "").strip() for c in rows[hi]]
        qcol = hdr.index("Question")
        wave_cols = [(j, int(h)) for j, h in enumerate(hdr) if h.isdigit()]

        def cell(r, col):  # a named column's text, or "" if absent
            j = hdr.index(col) if col in hdr else len(r)
            return str(r[j] or "").strip() if j < len(r) else ""

        for r in rows[hi + 1:]:
            if not r or qcol >= len(r) or not r[qcol]:
                continue
            asked = [w for j, w in wave_cols if j < len(r) and str(r[j] or "").strip()]
            out.append({"section": cell(r, "Section"), "subsection": cell(r, "Sub-Section"),
                        "question": generic_question(r[qcol]), "rotation": cell(r, "Rotation"),
                        "waves": asked})
    return out


def last_asked(log: list[dict], question: str) -> int | None:
    """The last wave in which a question containing `question` was asked."""
    def loose(t):  # the log and the sheets differ in commas and quotes, so drop punctuation
        return re.sub(r"[^a-z0-9\[\] ]", "", norm(t))

    q = loose(question)
    waves = [max(e["waves"]) for e in log if e["waves"] and q in loose(e["question"])]
    return max(waves) if waves else None


# ---------------------------------------------------------------- the store

class Store:
    """data/bics/: harvested questions, their catalogue, wave dates, and the question log."""

    def __init__(self, root: Path):
        self.root = root
        self.hdir = root / "harvest"
        self.hdir.mkdir(parents=True, exist_ok=True)
        self.state_path = root / "state.json"
        self.state = json.loads(self.state_path.read_text()) if self.state_path.exists() else {}
        self.state.setdefault("harvested", [])
        self.catalogue = self._read_csv(root / "catalogue.csv", key="id")
        self.waves = {int(k): v for k, v in self._read_csv(root / "waves.csv", key="wave").items()}
        self.log: list[dict] = []
        if (root / "question_log.csv").exists():
            with (root / "question_log.csv").open(newline="") as f:
                for r in csv.DictReader(f):
                    r["waves"] = [int(w) for w in r["waves"].split()]
                    self.log.append(r)

    @staticmethod
    def _read_csv(path: Path, key: str) -> dict:
        """CSV rows keyed by one column ({} if the file doesn't exist yet)."""
        if not path.exists():
            return {}
        with path.open(newline="") as f:
            return {r[key]: r for r in csv.DictReader(f)}

    def load(self, qid: str) -> list[tuple[int, str, str, float]]:
        """Harvest files are wide (wave, breakdown, one column per answer); return long rows."""
        path = self.hdir / f"{qid}.csv"
        if not path.exists():
            return []
        out = []
        with path.open(newline="") as f:
            for r in csv.DictReader(f):
                w, b = int(r.pop("wave")), r.pop("breakdown")
                out.extend((w, b, a, float(v)) for a, v in r.items() if v != "")
        return out

    def _write(self, qid: str, rows: list[tuple[int, str, str, float]], answers: list[str]) -> None:
        """Write one question's harvest file, wide: wave, breakdown, one column per answer
        (current answers first, then any older wordings)."""
        cells: dict[tuple[int, str], dict[str, float]] = {}
        for w, b, a, v in rows:
            cells.setdefault((w, b), {})[a] = v
        cols = answers + sorted({r[2] for r in rows} - set(answers))
        with (self.hdir / f"{qid}.csv").open("w", newline="") as f:
            wr = csv.writer(f)
            wr.writerow(["wave", "breakdown"] + cols)
            for (w, b) in sorted(cells):
                wr.writerow([w, b] + [cells[(w, b)].get(a, "") for a in cols])

    def add_wave(self, wave: int, path: Path, with_log: bool = False) -> int:
        """Merge one wave file into the store.

        For waves both files cover, the newer file wins (it carries ONS revisions); an
        older file only fills in waves the stored copy lacks. So files can be added in
        any order. `with_log` also refreshes the question log (pass it for the newest file).
        """
        tables, periods = read_wave(path)
        for w, p in periods.items():
            self.waves[w] = {"wave": w, **p}
        for t in tables:
            qid = question_id(t["question"], t["population"], t["weighting"])
            old = self.catalogue.get(qid, {})
            existing = self.load(qid)
            newest = wave >= int(old.get("source_wave", 0))  # is this file newer than the stored copy?
            if newest:  # it replaces every wave it covers
                new_waves = {r[0] for r in t["rows"]}
                rows = [r for r in existing if r[0] not in new_waves] + t["rows"]
            else:       # it only fills in waves the stored copy lacks
                have = {r[0] for r in existing}
                rows = existing + [r for r in t["rows"] if r[0] not in have]
            self._write(qid, rows, t["answers"])
            all_waves = sorted({r[0] for r in rows})
            # answers: union over time, newest wording first, so renamed options stay visible
            old_ans = [a for a in old.get("answers", "").split(" | ") if a]
            answers = list(dict.fromkeys(t["answers"] + old_ans if newest else old_ans + t["answers"]))
            self.catalogue[qid] = {
                "id": qid, "question": t["question"] if newest else old["question"],
                "population": t["population"], "weighting": t["weighting"],
                "answers": " | ".join(answers),
                "first_wave": all_waves[0], "last_wave": all_waves[-1], "n_waves": len(all_waves),
                "source_wave": max(wave, int(old.get("source_wave", 0))),
                "sheet": t["sheet"] if newest else old["sheet"],
            }
        if with_log:
            self.log = read_question_log(path)
        self.state["harvested"] = sorted(set(self.state["harvested"]) | {wave})
        return len(tables)

    def save(self) -> None:
        """Write the catalogue, wave dates, question log and state."""
        cols = ["id", "question", "population", "weighting", "answers",
                "first_wave", "last_wave", "n_waves", "source_wave", "sheet"]
        with (self.root / "catalogue.csv").open("w", newline="") as f:
            wr = csv.DictWriter(f, fieldnames=cols)
            wr.writeheader()
            for row in sorted(self.catalogue.values(), key=lambda r: (-int(r["last_wave"]), r["question"])):
                wr.writerow({c: row[c] for c in cols})
        with (self.root / "waves.csv").open("w", newline="") as f:
            wr = csv.DictWriter(f, fieldnames=["wave", "ref_start", "ref_end", "live_start", "live_end"])
            wr.writeheader()
            for w in sorted(self.waves):
                wr.writerow({"wave": w, **{k: self.waves[w].get(k) for k in wr.fieldnames[1:]}})
        with (self.root / "question_log.csv").open("w", newline="") as f:
            wr = csv.writer(f)
            wr.writerow(["section", "subsection", "question", "rotation",
                         "first_wave", "last_wave", "n_waves", "waves"])
            for e in sorted(self.log, key=lambda e: -(max(e["waves"]) if e["waves"] else 0)):
                ws = e["waves"]
                wr.writerow([e["section"], e["subsection"], e["question"], e["rotation"],
                             min(ws) if ws else "", max(ws) if ws else "", len(ws), " ".join(map(str, ws))])
        self.state_path.write_text(json.dumps(self.state, indent=1) + "\n")


def fetch_waves(store: Store, wanted: set[int], index: dict[int, str], local_dir: Path | None = None) -> list[int]:
    """Harvest each wanted wave not already in the store: from `local_dir` (files named
    w<N>.xlsx or w<NNN>.xlsx) if given and present, otherwise downloaded from ONS.
    The newest published wave also refreshes the question log."""
    todo = sorted(w for w in wanted if w not in store.state["harvested"])
    latest = max(index)
    with tempfile.TemporaryDirectory() as tmp:
        for w in todo:
            local = [p for p in (local_dir / f"w{w}.xlsx", local_dir / f"w{w:03d}.xlsx")
                     if p.exists()] if local_dir else []
            path = local[0] if local else download_wave(index[w], Path(tmp) / f"w{w}.xlsx")
            n = store.add_wave(w, path, with_log=(w == latest))
            print(f"[bics] wave {w}: {n} time-series tables{' (+ question log)' if w == latest else ''}")
    return todo
